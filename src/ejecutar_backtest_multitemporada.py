#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Backtest 1X2 temporal auditable. Requisitos: pandas numpy requests openpyxl.

No calcula ROI/EV porque no descarga cuotas históricas verificadas. El resultado
es exploratorio y no prueba rentabilidad futura.
"""
from pathlib import Path
from datetime import datetime, timezone
import hashlib, json, math
import requests
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment

BASE = Path(__file__).resolve().parent
RAW = BASE / "datos_futbol_raw"
RAW.mkdir(parents=True, exist_ok=True)
OUT = BASE / "backtest_multitemporada_resultados.xlsx"
MANIFEST = BASE / "manifest_fuentes.json"
LEAGUES = {
    "Premier League": "english-premier-league",
    "LaLiga": "spanish-la-liga",
    "Serie A": "italian-serie-a",
    "Bundesliga": "german-bundesliga",
    "Ligue 1": "french-ligue-1",
}
SEASONS = ["2324", "2425", "2526"]
URL = "https://datahub.io/football/{slug}/_r/-/season-{season}.csv"
RUN_UTC = datetime.now(timezone.utc).isoformat(timespec="seconds")

def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def load_source(league, slug, season):
    url = URL.format(slug=slug, season=season)
    path = RAW / f"{slug}_season-{season}.csv"
    if not path.exists() or path.stat().st_size == 0:
        r = requests.get(url, timeout=40)
        r.raise_for_status()
        if len(r.content) < 500:
            raise ValueError(f"Respuesta demasiado pequeña: {url}")
        path.write_bytes(r.content)
    digest = sha256(path)
    src = pd.read_csv(path)
    required = {"Date", "HomeTeam", "AwayTeam", "FTHG", "FTAG"}
    if not required.issubset(src.columns):
        raise ValueError(f"Faltan columnas: {sorted(required - set(src.columns))}")
    raw_date = src["Date"].astype("string").str.strip()
    dates = pd.to_datetime(raw_date, format="%Y-%m-%d", errors="coerce")
    alt = dates.isna() & raw_date.str.match(r"^\d{1,2}/\d{1,2}/\d{4}$", na=False)
    dates.loc[alt] = pd.to_datetime(raw_date.loc[alt], format="%d/%m/%Y", errors="coerce")
    out = src.rename(columns={"Date":"date_raw","HomeTeam":"home","AwayTeam":"away","FTHG":"hg","FTAG":"ag"}).copy()
    out["date"] = dates
    out["hg"] = pd.to_numeric(out["hg"], errors="coerce")
    out["ag"] = pd.to_numeric(out["ag"], errors="coerce")
    for c in ["home","away"]:
        out[c] = out[c].astype("string").str.strip()
    out["league"], out["season"] = league, season
    out["source_url"], out["source_file"], out["source_sha256"] = url, path.name, digest
    out["downloaded_utc"] = RUN_UTC
    out["source_license"] = "ODC-PDDL-1.0 según ficha DataHub; verificar términos vigentes"
    manifest = {"league":league,"season":season,"url":url,"file":path.name,"sha256":digest,
                "bytes":path.stat().st_size,"rows":len(out),"downloaded_utc":RUN_UTC,"status":"OK"}
    return out[["league","season","date","date_raw","home","away","hg","ag","source_url","source_file","source_sha256","downloaded_utc","source_license"]], manifest

def clean_data(raw):
    d = raw.copy()
    bad = d[d.date.isna() | d.home.isna() | d.away.isna() | d.home.eq("") | d.away.eq("") |
            d.hg.isna() | d.ag.isna() | (d.hg < 0) | (d.ag < 0)].copy()
    d = d.drop(index=bad.index)
    nonint = d[(d.hg % 1 != 0) | (d.ag % 1 != 0)].copy()
    d = d.drop(index=nonint.index)
    d["hg"], d["ag"] = d.hg.astype(int), d.ag.astype(int)
    key = ["league","date","home","away"]
    exact = key + ["hg","ag"]
    duplicates = d[d.duplicated(exact, keep="first")].copy()
    d = d.drop_duplicates(exact, keep="first").copy()
    counts = d.groupby(key, dropna=False)[["hg","ag"]].apply(lambda x: len(x.drop_duplicates()))
    conflict_keys = counts[counts > 1].index
    if len(conflict_keys):
        idx = pd.MultiIndex.from_frame(d[key])
        bad_idx = pd.MultiIndex.from_tuples(list(conflict_keys), names=key)
        conflicts = d[idx.isin(bad_idx)].copy()
        d = d.loc[~idx.isin(bad_idx)].copy()
    else:
        conflicts = d.iloc[0:0].copy()
    d = d.sort_values(["league","date","season","home","away"]).reset_index(drop=True)
    return d, bad, nonint, duplicates, conflicts

def pois(lam, max_goals=10):
    vals = [math.exp(-lam)]
    for k in range(1, max_goals+1):
        vals.append(vals[-1] * lam / k)
    a = np.array(vals, dtype=float)
    return a / a.sum()

def probs_1x2(lh, la):
    m = np.outer(pois(lh), pois(la))  # filas: goles local; columnas: visitante
    p = np.array([np.tril(m,-1).sum(), np.trace(m), np.triu(m,1).sum()], dtype=float)
    return p / p.sum()

def form(history, team, venue=None):
    h = [x for x in history if x["team"] == team and (venue is None or x["venue"] == venue)][-5:]
    if not h: return None
    return (float(np.mean([x["gf"] for x in h])), float(np.mean([x["ga"] for x in h])), len(h))

def events(r):
    return [{"team":r.home,"venue":"home","gf":int(r.hg),"ga":int(r.ag),"date":r.date},
            {"team":r.away,"venue":"away","gf":int(r.ag),"ga":int(r.hg),"date":r.date}]

def evaluate(data):
    detail, skipped = [], []
    for league, group in data.groupby("league", sort=False):
        history = []
        group = group.sort_values(["date","season","home","away"])
        for r in group.itertuples(index=False):
            prior = [e for e in history if e["date"] < r.date] # excluye partidos del mismo día
            h, a = form(prior,r.home), form(prior,r.away)
            if h is None or a is None or h[2] < 5 or a[2] < 5:
                skipped.append({"league":league,"season":r.season,"date":r.date,"home":r.home,"away":r.away,
                                "reason":"Menos de 5 partidos previos para uno o ambos equipos"})
                history.extend(events(r)); continue
            hvenue, aVenue = form(prior,r.home,"home"), form(prior,r.away,"away")
            uh = hvenue if hvenue and hvenue[2] >= 3 else h
            ua = aVenue if aVenue and aVenue[2] >= 3 else a
            base_h, base_a = (h[0]+a[1])/2, (a[0]+h[1])/2
            models = {"multipliers_1.10_0.90":(max(.05,1.10*base_h),max(.05,.90*base_a)),
                      "neutral_1.00_1.00":(max(.05,base_h),max(.05,base_a)),
                      "venue_adjusted":(max(.05,(uh[0]+ua[1])/2),max(.05,(ua[0]+uh[1])/2))}
            actual = "1" if r.hg>r.ag else ("X" if r.hg==r.ag else "2")
            rec = {"league":league,"season":r.season,"date":r.date,"home":r.home,"away":r.away,
                   "hg":int(r.hg),"ag":int(r.ag),"actual_1X2":actual,"source_url":r.source_url,
                   "source_sha256":r.source_sha256,"prior_home_matches":h[2],"prior_away_matches":a[2]}
            labels = ["1","X","2"]
            y = np.array([float(actual==x) for x in labels])
            for name,(lh,la) in models.items():
                p = probs_1x2(lh,la); pred = labels[int(np.argmax(p))]
                rec.update({f"{name}_lambda_home":lh,f"{name}_lambda_away":la,
                            f"{name}_p1":p[0],f"{name}_px":p[1],f"{name}_p2":p[2],
                            f"{name}_prediction":pred,f"{name}_correct":int(pred==actual),
                            f"{name}_brier":float(np.sum((p-y)**2)),
                            f"{name}_logloss":float(-math.log(max(p[labels.index(actual)],1e-15)))})
            detail.append(rec)
            history.extend(events(r))
    return pd.DataFrame(detail), pd.DataFrame(skipped)

def summarize(detail):
    names = ["multipliers_1.10_0.90","neutral_1.00_1.00","venue_adjusted"]
    global_rows, grouped = [], []
    if detail.empty: return pd.DataFrame(), pd.DataFrame()
    for name in names:
        global_rows.append({"Modelo":name,"Partidos":len(detail),"Aciertos":int(detail[f"{name}_correct"].sum()),
            "Exactitud":float(detail[f"{name}_correct"].mean()),"Brier":float(detail[f"{name}_brier"].mean()),
            "LogLoss":float(detail[f"{name}_logloss"].mean()),
            "Nota":"Walk-forward exploratorio; no mide ROI ni demuestra rentabilidad"})
    for (league,season), g in detail.groupby(["league","season"]):
        for name in names:
            grouped.append({"Liga":league,"Temporada":season,"Modelo":name,"Partidos":len(g),
                "Aciertos":int(g[f"{name}_correct"].sum()),"Exactitud":float(g[f"{name}_correct"].mean()),
                "Brier":float(g[f"{name}_brier"].mean()),"LogLoss":float(g[f"{name}_logloss"].mean())})
    return pd.DataFrame(global_rows), pd.DataFrame(grouped)

def main():
    frames, sources, errors = [], [], []
    for league, slug in LEAGUES.items():
        for season in SEASONS:
            try:
                frame, m = load_source(league,slug,season); frames.append(frame); sources.append(m)
            except Exception as e:
                errors.append({"league":league,"season":season,"url":URL.format(slug=slug,season=season),
                               "status":"ERROR","error":str(e)})
    if not frames: raise RuntimeError("No se descargaron CSV. Compruebe internet o coloque los 15 CSV en datos_futbol_raw/.")
    raw = pd.concat(frames,ignore_index=True)
    clean,bad,nonint,dups,conflicts = clean_data(raw)
    detail,skipped = evaluate(clean)
    metrics,by_season = summarize(detail)
    quality = pd.DataFrame([
        ["Ejecución UTC",RUN_UTC,"Trazabilidad"],["Fuentes OK",len(sources),f"Esperadas: {len(LEAGUES)*len(SEASONS)}"],
        ["Fuentes fallidas",len(errors),"Revisar manifiesto"],["Filas importadas",len(raw),"Antes de limpieza"],
        ["Filas inválidas",len(bad),"Fecha/equipo/marcador ausente o marcador negativo"],
        ["Marcadores no enteros",len(nonint),"Excluidos y registrados"],["Duplicados idénticos retirados",len(dups),"Se conserva una copia"],
        ["Conflictos de marcador excluidos",len(conflicts),"Misma liga/fecha/local/visitante, marcadores distintos"],
        ["Partidos únicos limpios",len(clean),"Después de limpieza explícita"],["Partidos evaluables",len(detail),"Mínimo 5 partidos previos por equipo"],
        ["Partidos omitidos",len(skipped),"Motivo por partido en Omitidos"],["Cuotas/EV/ROI","No calculados","No hay cuotas históricas verificadas"],
        ["Corte temporal","Solo partidos anteriores a la fecha objetivo","Los resultados del mismo día no son antecedentes"]
    ],columns=["Control","Resultado","Nota"])
    MANIFEST.write_text(json.dumps({"generated_utc":RUN_UTC,"sources":sources,"errors":errors},ensure_ascii=False,indent=2),encoding="utf-8")
    with pd.ExcelWriter(OUT,engine="openpyxl") as writer:
        pd.DataFrame(sources+errors).to_excel(writer,index=False,sheet_name="Fuentes_y_descargas")
        metrics.to_excel(writer,index=False,sheet_name="Metricas_globales")
        by_season.to_excel(writer,index=False,sheet_name="Metricas_liga_temporada")
        detail.to_excel(writer,index=False,sheet_name="Detalle_backtest")
        skipped.to_excel(writer,index=False,sheet_name="Omitidos")
        dups.to_excel(writer,index=False,sheet_name="Duplicados_exactos")
        conflicts.to_excel(writer,index=False,sheet_name="Conflictos_marcador")
        bad.to_excel(writer,index=False,sheet_name="Filas_invalidas")
        nonint.to_excel(writer,index=False,sheet_name="Marcadores_no_enteros")
        quality.to_excel(writer,index=False,sheet_name="Controles_calidad")
        raw.to_excel(writer,index=False,sheet_name="Datos_importados")
        for ws in writer.book.worksheets:
            ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
            for c in ws[1]:
                c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E78")
                c.alignment=Alignment(wrap_text=True,vertical="top")
            for col in ws.columns:
                letter=col[0].column_letter
                maxlen=max((len(str(c.value)) if c.value is not None else 0) for c in col[:100])
                ws.column_dimensions[letter].width=min(max(maxlen+2,12),48)
            for row in ws.iter_rows(min_row=2):
                for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)
    print(f"Resultado: {OUT}\nManifiesto: {MANIFEST}")
    print(f"Fuentes OK {len(sources)}/15; filas {len(raw)}; limpias {len(clean)}; evaluables {len(detail)}; omitidas {len(skipped)}")
    print(f"Duplicados retirados {len(dups)}; conflictos excluidos {len(conflicts)}")
    if errors: print("ADVERTENCIA: ejecución parcial; no tratarla como backtest completo.")

if __name__ == "__main__":
    main()
