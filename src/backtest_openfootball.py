#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Reproduce backtest 1X2 from OpenFootball JSON. No odds/ROI."""
from pathlib import Path
from datetime import datetime, timezone
import hashlib, json, math
import requests
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment

BASE=Path(__file__).resolve().parent
RAW=BASE/"data"/"openfootball"
RAW.mkdir(parents=True,exist_ok=True)
OUT=BASE/"backtest_openfootball_resultados.xlsx"
MANIFEST=BASE/"manifest_openfootball.json"
SEASONS=["2023-24","2024-25","2025-26"]
LEAGUES={"Premier League":"en.1","LaLiga":"es.1","Serie A":"it.1","Bundesliga":"de.1","Ligue 1":"fr.1"}
STAMP=datetime.now(timezone.utc).isoformat(timespec="seconds")

def sha256(path):
    h=hashlib.sha256()
    with open(path,"rb") as f:
        for b in iter(lambda:f.read(1<<20),b""): h.update(b)
    return h.hexdigest()

def fetch(league,code,season):
    url=f"https://raw.githubusercontent.com/openfootball/football.json/master/{season}/{code}.json"
    path=RAW/f"{season}_{code}.json"
    if not path.exists():
        r=requests.get(url,timeout=45); r.raise_for_status()
        path.write_bytes(r.content)
    obj=json.loads(path.read_text(encoding="utf-8"))
    # OpenFootball has used both an object with a "matches" array and a
    # top-level array. Normalize both shapes before parsing.
    if isinstance(obj, dict):
        matches=obj.get("matches", [])
        source_name=obj.get("name")
    elif isinstance(obj, list):
        matches=obj
        source_name=f"{league} {season}"
    else:
        raise ValueError(f"Formato JSON inesperado: {type(obj).__name__}")
    if not isinstance(matches, list):
        raise ValueError("El campo matches no es una lista")
    rows=[]
    for m in matches:
        if not isinstance(m, dict):
            continue
        score_obj=m.get("score",{})
        # Scores may be {"ft":[home,away], "ht":[...]} or a direct
        # [home, away] pair. Treat either shape consistently.
        if isinstance(score_obj, dict):
            score=score_obj.get("ft")
        elif isinstance(score_obj, list):
            score=score_obj
        else:
            score=None
        hg=score[0] if isinstance(score,list) and len(score)==2 else None
        ag=score[1] if isinstance(score,list) and len(score)==2 else None
        rows.append({"league":league,"season":season,"date":m.get("date"),
          "home":m.get("team1"),"away":m.get("team2"),"hg":hg,"ag":ag,
          "round":m.get("round"),"source_url":url,"source_file":path.name,"source_sha256":sha256(path)})
    return rows,{"league":league,"season":season,"url":url,"file":path.name,
      "sha256":sha256(path),"bytes":path.stat().st_size,"rows":len(rows),"source_name":source_name,
      "downloaded_utc":STAMP,"status":"OK"}

def poisson(lam,n=12):
    a=[math.exp(-lam)]
    for k in range(1,n+1): a.append(a[-1]*lam/k)
    s=sum(a); return [x/s for x in a]

def p1x2(lh,la):
    h,a=poisson(lh),poisson(la); p=[0.0,0.0,0.0]
    for i,x in enumerate(h):
        for j,y in enumerate(a):
            p[0 if i>j else 1 if i==j else 2]+=x*y
    z=sum(p); return [x/z for x in p]

def form(hist,team,venue=None):
    x=[r for r in hist if r["team"]==team and (venue is None or r["venue"]==venue)][-5:]
    if not x:return None
    return (sum(r["gf"] for r in x)/len(x),sum(r["ga"] for r in x)/len(x),len(x))

def clean(rows):
    valid=[]; invalid=[]
    for r in rows:
        try:
            date=pd.to_datetime(r["date"],format="%Y-%m-%d",errors="raise")
            hg,ag=int(r["hg"]),int(r["ag"])
            if hg<0 or ag<0 or not r["home"] or not r["away"]: raise ValueError()
            valid.append({**r,"date":date,"hg":hg,"ag":ag})
        except Exception: invalid.append(r)
    key=lambda r:(r["league"],r["date"],r["home"],r["away"])
    scorekey=lambda r:(*key(r),r["hg"],r["ag"])
    seen=set(); unique=[]; dups=[]
    for r in valid:
        k=scorekey(r)
        if k in seen: dups.append(r)
        else: seen.add(k); unique.append(r)
    scores={}
    for r in unique:scores.setdefault(key(r),set()).add((r["hg"],r["ag"]))
    badkeys={k for k,v in scores.items() if len(v)>1}
    conflicts=[r for r in unique if key(r) in badkeys]
    clean=[r for r in unique if key(r) not in badkeys]
    clean.sort(key=lambda r:(r["league"],r["date"],r["home"],r["away"]))
    return clean,invalid,dups,conflicts

def backtest(rows):
    detail=[]; skipped=[]
    models=["multipliers_1.10_0.90","neutral_1.00_1.00","venue_adjusted"]
    for league in LEAGUES:
        hist=[]
        for r in [x for x in rows if x["league"]==league]:
            prior=[x for x in hist if x["date"]<r["date"]]
            h,a=form(prior,r["home"]),form(prior,r["away"])
            if not h or not a or h[2]<5 or a[2]<5:
                skipped.append({**r,"reason":"Menos de 5 partidos previos para uno o ambos equipos"})
            else:
                hh,aa=form(prior,r["home"],"home"),form(prior,r["away"],"away")
                uh=hh if hh and hh[2]>=3 else h; ua=aa if aa and aa[2]>=3 else a
                bh=(h[0]+a[1])/2; ba=(a[0]+h[1])/2
                lambdas={"multipliers_1.10_0.90":(max(.05,1.1*bh),max(.05,.9*ba)),
                  "neutral_1.00_1.00":(max(.05,bh),max(.05,ba)),
                  "venue_adjusted":(max(.05,(uh[0]+ua[1])/2),max(.05,(ua[0]+uh[1])/2))}
                actual="1" if r["hg"]>r["ag"] else "X" if r["hg"]==r["ag"] else "2"
                rec={k:r[k] for k in ["league","season","date","home","away","hg","ag","source_url","source_sha256"]}
                rec["actual_1X2"]=actual
                # Baseline de frecuencia histórica con suavizado Laplace y solo partidos anteriores.
                prior_matches=[x for x in rows if x["league"]==league and x["date"]<r["date"] and x["hg"] is not None and x["ag"] is not None]
                outcomes=["1" if x["hg"]>x["ag"] else "X" if x["hg"]==x["ag"] else "2" for x in prior_matches]
                counts=[outcomes.count(label) for label in ["1","X","2"]]
                probs_base=[(c+1)/(len(outcomes)+3) for c in counts]
                idx=["1","X","2"].index(actual)
                pred_base=["1","X","2"][probs_base.index(max(probs_base))]
                rec.update({"baseline_frequency_prior_n":len(outcomes),
                    "baseline_frequency_p1":probs_base[0],"baseline_frequency_px":probs_base[1],
                    "baseline_frequency_p2":probs_base[2],"baseline_frequency_prediction":pred_base,
                    "baseline_frequency_correct":int(pred_base==actual),
                    "baseline_frequency_brier":sum((v-(i==idx))**2 for i,v in enumerate(probs_base)),
                    "baseline_frequency_logloss":-math.log(max(probs_base[idx],1e-15))})
                for m,(lh,la) in lambdas.items():
                    p=p1x2(lh,la); labels=["1","X","2"]; idx=labels.index(actual); pred=labels[p.index(max(p))]
                    rec.update({f"{m}_lambda_home":lh,f"{m}_lambda_away":la,
                      f"{m}_p1":p[0],f"{m}_px":p[1],f"{m}_p2":p[2],f"{m}_prediction":pred,
                      f"{m}_correct":int(pred==actual),f"{m}_brier":sum((v-(i==idx))**2 for i,v in enumerate(p)),
                      f"{m}_logloss":-math.log(max(p[idx],1e-15))})
                detail.append(rec)
            hist.extend([{"team":r["home"],"venue":"home","gf":r["hg"],"ga":r["ag"],"date":r["date"]},
                         {"team":r["away"],"venue":"away","gf":r["ag"],"ga":r["hg"],"date":r["date"]}])
    return pd.DataFrame(detail),pd.DataFrame(skipped)


def paired_bootstrap_comparisons(detail, n_boot=3000, seed=20261014):
    """Bootstrap pareado por partido; diferencia = métrica A menos métrica B."""
    comparisons = [
        ("multipliers_1.10_0.90", "baseline_frequency"),
        ("venue_adjusted", "baseline_frequency"),
        ("multipliers_1.10_0.90", "venue_adjusted"),
        ("multipliers_1.10_0.90", "neutral_1.00_1.00"),
    ]
    metric_columns = {"Exactitud": ("correct", True), "Brier": ("brier", False), "LogLoss": ("logloss", False)}
    columns = ["Modelo_A", "Modelo_B", "Metrica", "N", "Diferencia_A_menos_B",
               "IC95_inferior", "IC95_superior", "Remuestras", "Semilla", "Metodo", "Interpretacion"]
    if detail.empty:
        return pd.DataFrame(columns=columns)
    rng = np.random.default_rng(seed)
    n = len(detail)
    results = []
    for model_a, model_b in comparisons:
        for metric, (suffix, higher_is_better) in metric_columns.items():
            delta = (detail[f"{model_a}_{suffix}"].to_numpy(dtype=float)
                     - detail[f"{model_b}_{suffix}"].to_numpy(dtype=float))
            boot = np.empty(n_boot, dtype=float)
            for i in range(n_boot):
                boot[i] = delta[rng.integers(0, n, size=n)].mean()
            estimate = float(delta.mean())
            lo, hi = float(np.quantile(boot, .025)), float(np.quantile(boot, .975))
            if lo > 0:
                verdict = "A mejor" if higher_is_better else "B mejor"
            elif hi < 0:
                verdict = "A mejor" if not higher_is_better else "B mejor"
            else:
                verdict = "IC95 incluye cero; diferencia no concluyente"
            results.append({"Modelo_A": model_a, "Modelo_B": model_b, "Metrica": metric, "N": n,
                "Diferencia_A_menos_B": estimate, "IC95_inferior": lo, "IC95_superior": hi,
                "Remuestras": n_boot, "Semilla": seed,
                "Metodo": "Bootstrap pareado por partido; mismas filas para ambos modelos",
                "Interpretacion": verdict})
    return pd.DataFrame(results)

def main():
    rows=[]; manifests=[]; errors=[]
    for league,code in LEAGUES.items():
        for season in SEASONS:
            try:
                rr,mm=fetch(league,code,season); rows.extend(rr); manifests.append(mm)
            except Exception as e: errors.append({"league":league,"season":season,"status":"ERROR","error":str(e)})
    if not rows: raise RuntimeError("No se pudo recuperar ninguna fuente OpenFootball.")
    clean_rows,invalid,dups,conflicts=clean(rows)
    detail,skipped=backtest(clean_rows)
    models=["multipliers_1.10_0.90","neutral_1.00_1.00","venue_adjusted","baseline_frequency"]
    glob=[]; by=[]; metrics_ci=[]; calibration=[]
    rng=np.random.default_rng(20261014); n=len(detail)
    for m in models:
        if n:
            glob.append({"Modelo":m,"Partidos":n,"Exactitud":detail[f"{m}_correct"].mean(),
              "Brier":detail[f"{m}_brier"].mean(),"LogLoss":detail[f"{m}_logloss"].mean(),
              "Nota":"Frecuencia histórica suavizada; solo partidos anteriores" if m=="baseline_frequency" else "Sin cuotas, EV ni ROI"})
            for metric in ["correct","brier","logloss"]:
                vals=detail[f"{m}_{metric}"].to_numpy(dtype=float)
                boot=np.array([rng.choice(vals,size=n,replace=True).mean() for _ in range(1000)])
                metrics_ci.append({"Modelo":m,"Metrica":{"correct":"Exactitud","brier":"Brier","logloss":"LogLoss"}[metric],
                    "N":n,"Estimacion":float(vals.mean()),"IC95_inferior":float(np.quantile(boot,.025)),
                    "IC95_superior":float(np.quantile(boot,.975)),"Metodo":"Bootstrap por partido, 1000 remuestras, semilla fija"})
            probs=detail[[f"{m}_p1",f"{m}_px",f"{m}_p2"]].to_numpy(dtype=float)
            actuals=detail["actual_1X2"].to_numpy(); labels=["1","X","2"]; edges=np.linspace(0,1,11)
            for j,label in enumerate(labels):
                pp=probs[:,j]; yy=(actuals==label).astype(float)
                for bi in range(10):
                    mask=(pp>=edges[bi]) & ((pp<edges[bi+1]) if bi<9 else (pp<=edges[bi+1]))
                    if mask.any():
                        calibration.append({"Modelo":m,"Clase":label,"Banda":f"{edges[bi]:.1f}-{edges[bi+1]:.1f}",
                            "N":int(mask.sum()),"Probabilidad_media":float(pp[mask].mean()),
                            "Frecuencia_observada":float(yy[mask].mean()),
                            "Error_absoluto_calibracion":float(abs(pp[mask].mean()-yy[mask].mean()))})
    for (league,season),g in detail.groupby(["league","season"]):
        for m in models:
            by.append({"Liga":league,"Temporada":season,"Modelo":m,"N":len(g),
              "Exactitud":g[f"{m}_correct"].mean(),"Brier":g[f"{m}_brier"].mean(),"LogLoss":g[f"{m}_logloss"].mean()})
    paired_comparisons = paired_bootstrap_comparisons(detail, n_boot=3000, seed=20261014)
    quality=pd.DataFrame([
      ["Fuente", "OpenFootball football.json","Fuente efectivamente consultada"],
      ["Filas fuente",len(rows),"15 JSON"],
      ["Marcadores completos",len(rows)-len(invalid),"Antes de deduplicar"],
      ["Filas inválidas/incompletas",len(invalid),"Revisar hoja correspondiente"],
      ["Duplicados exactos retirados",len(dups),"Conserva una copia"],
      ["Filas en conflicto excluidas",len(conflicts),"Misma liga/fecha/local/visitante con marcadores distintos"],
      ["Partidos limpios",len(clean_rows),"Después de limpieza"],
      ["Evaluables",len(detail),"Al menos 5 partidos previos por equipo"],
      ["Omitidos",len(skipped),"Detalle por partido"],
      ["EV/ROI","No calculados","No hay cuotas históricas verificadas"],
      ["Corte temporal","Solo partidos con fecha anterior","Sin resultados del mismo día en el historial"]],columns=["Control","Resultado","Nota"])
    MANIFEST.write_text(json.dumps({"run_utc":STAMP,"sources":manifests,"errors":errors},ensure_ascii=False,indent=2),encoding="utf-8")
    with pd.ExcelWriter(OUT,engine="openpyxl") as w:
        pd.DataFrame(manifests+errors).to_excel(w,index=False,sheet_name="Fuentes")
        pd.DataFrame(glob).to_excel(w,index=False,sheet_name="Metricas_globales")
        pd.DataFrame(by).to_excel(w,index=False,sheet_name="Metricas_liga_temporada")
        pd.DataFrame(metrics_ci).to_excel(w,index=False,sheet_name="Intervalos_confianza")
        pd.DataFrame(calibration).to_excel(w,index=False,sheet_name="Calibracion")
        paired_comparisons.to_excel(w,index=False,sheet_name="Comparaciones_pareadas")
        detail.to_excel(w,index=False,sheet_name="Detalle_backtest")
        skipped.to_excel(w,index=False,sheet_name="Omitidos")
        pd.DataFrame(dups).to_excel(w,index=False,sheet_name="Duplicados")
        pd.DataFrame(conflicts).to_excel(w,index=False,sheet_name="Conflictos")
        pd.DataFrame(invalid).to_excel(w,index=False,sheet_name="Invalidos")
        pd.DataFrame(rows).to_excel(w,index=False,sheet_name="Datos_fuente")
        quality.to_excel(w,index=False,sheet_name="Controles")
        for ws in w.book.worksheets:
            ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
            for c in ws[1]:
                c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E78")
                c.alignment=Alignment(wrap_text=True,vertical="top")
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width=min(max(max(len(str(c.value or "")) for c in col[:100])+2,12),48)
    print(f"Fuentes OK: {len(manifests)}/15; errores: {len(errors)}; filas fuente: {len(rows)}")
    print(f"Completos: {len(rows)-len(invalid)}; limpios: {len(clean_rows)}; evaluables: {len(detail)}; omitidos: {len(skipped)}")
    print(f"Excel: {OUT}; manifiesto: {MANIFEST}")

if __name__=="__main__": main()
