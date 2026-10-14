#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Reproduce un backtest temporal 1X2 con datos JSON de OpenFootball.

Instalación: pip install pandas numpy requests openpyxl
Ejecución: python src/ejecutar_backtest_openfootball.py

No usa cuotas. Sus métricas no demuestran ROI ni rentabilidad.
"""
from pathlib import Path
from datetime import datetime, timezone
import hashlib, json, math
import requests
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "openfootball_raw"
RAW.mkdir(parents=True, exist_ok=True)
OUT = ROOT / "reports" / "backtest_openfootball.xlsx"
MANIFEST = ROOT / "reports" / "manifest_openfootball.json"
SEASONS = ["2023-24", "2024-25", "2025-26"]
LEAGUES = {
    "Premier League": "en.1",
    "LaLiga": "es.1",
    "Serie A": "it.1",
    "Bundesliga": "de.1",
    "Ligue 1": "fr.1",
}
BASE_URL = "https://raw.githubusercontent.com/openfootball/football.json/master/{season}/{code}.json"
RUN_UTC = datetime.now(timezone.utc).isoformat(timespec="seconds")

def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()

def fetch(league, code, season):
    url = BASE_URL.format(season=season, code=code)
    path = RAW / f"{season}_{code}.json"
    if not path.exists():
        response = requests.get(url, timeout=40)
        response.raise_for_status()
        payload = response.content
        if len(payload) < 100:
            raise ValueError(f"Archivo inesperadamente pequeño: {url}")
        path.write_bytes(payload)
    digest = sha256(path)
    obj = json.loads(path.read_text(encoding="utf-8"))
    matches = obj.get("matches")
    if not isinstance(matches, list):
        raise ValueError(f"No existe una lista 'matches': {url}")
    rows = []
    for m in matches:
        score = (m.get("score") or {}).get("ft")
        hg = score[0] if isinstance(score, list) and len(score) == 2 else None
        ag = score[1] if isinstance(score, list) and len(score) == 2 else None
        rows.append({
            "league": league, "season": season, "date": m.get("date"),
            "home": m.get("team1"), "away": m.get("team2"),
            "hg": hg, "ag": ag, "round": m.get("round"),
            "source_url": url, "source_sha256": digest, "source_file": path.name,
        })
    manifest = {"league": league, "season": season, "url": url, "file": path.name,
                "sha256": digest, "bytes": path.stat().st_size, "rows": len(rows),
                "complete_scores": sum(r["hg"] is not None and r["ag"] is not None for r in rows),
                "downloaded_utc": RUN_UTC, "status": "OK"}
    return rows, manifest

def poisson(lam, max_goals=10):
    values = [math.exp(-lam)]
    for k in range(1, max_goals + 1):
        values.append(values[-1] * lam / k)
    total = sum(values)
    return [v / total for v in values]

def probs_1x2(lh, la):
    home, away = poisson(lh), poisson(la)
    ph = px = pa = 0.0
    for i, p_home in enumerate(home):
        for j, p_away in enumerate(away):
            value = p_home * p_away
            if i > j: ph += value
            elif i == j: px += value
            else: pa += value
    total = ph + px + pa
    return [ph/total, px/total, pa/total]

def form(history, team, venue=None):
    rows = [e for e in history if e["team"] == team and (venue is None or e["venue"] == venue)][-5:]
    if not rows: return None
    return (sum(e["gf"] for e in rows)/len(rows), sum(e["ga"] for e in rows)/len(rows), len(rows))

def events(row):
    return [
        {"team":row["home"],"venue":"home","gf":int(row["hg"]),"ga":int(row["ag"]),"date":row["date"]},
        {"team":row["away"],"venue":"away","gf":int(row["ag"]),"ga":int(row["hg"]),"date":row["date"]},
    ]

def clean(raw):
    d = raw.copy()
    d["date"] = pd.to_datetime(d["date"], format="%Y-%m-%d", errors="coerce")
    for c in ["league","season","home","away"]:
        d[c] = d[c].astype("string").str.strip()
    d["hg"] = pd.to_numeric(d["hg"], errors="coerce")
    d["ag"] = pd.to_numeric(d["ag"], errors="coerce")
    invalid = d[d["date"].isna() | d["home"].isna() | d["away"].isna() |
                d["home"].eq("") | d["away"].eq("") | d["hg"].isna() | d["ag"].isna() |
                (d["hg"] < 0) | (d["ag"] < 0)].copy()
    valid = d.drop(index=invalid.index).copy()
    nonint = valid[(valid["hg"] % 1 != 0) | (valid["ag"] % 1 != 0)].copy()
    valid = valid.drop(index=nonint.index)
    valid["hg"] = valid["hg"].astype(int)
    valid["ag"] = valid["ag"].astype(int)
    fixture = ["league","date","home","away"]
    exact = fixture + ["hg","ag"]
    dup = valid[valid.duplicated(exact, keep="first")].copy()
    unique = valid.drop_duplicates(exact, keep="first").copy()
    score_count = unique.groupby(fixture, dropna=False)[["hg","ag"]].apply(lambda x: len(x.drop_duplicates()))
    bad_keys = score_count[score_count > 1].index
    if len(bad_keys):
        idx = pd.MultiIndex.from_frame(unique[fixture])
        bad_idx = pd.MultiIndex.from_tuples(list(bad_keys), names=fixture)
        conflicts = unique.loc[idx.isin(bad_idx)].copy()
        unique = unique.loc[~idx.isin(bad_idx)].copy()
    else:
        conflicts = unique.iloc[0:0].copy()
    return unique.sort_values(["league","date","home","away"]).reset_index(drop=True), invalid, nonint, dup, conflicts

def backtest(data):
    detail, skipped = [], []
    for league, group in data.groupby("league", sort=False):
        history = []
        for r in group.sort_values(["date","home","away"]).to_dict("records"):
            prior = [e for e in history if e["date"] < r["date"]]
            h, a = form(prior,r["home"]), form(prior,r["away"])
            if not h or not a or h[2] < 5 or a[2] < 5:
                skipped.append({**r,"reason":"Menos de 5 partidos previos para uno o ambos equipos"})
                history.extend(events(r))
                continue
            hh, aa = form(prior,r["home"],"home"), form(prior,r["away"],"away")
            uh = hh if hh and hh[2] >= 3 else h
            ua = aa if aa and aa[2] >= 3 else a
            bh, ba = (h[0]+a[1])/2, (a[0]+h[1])/2
            actual = "1" if r["hg"] > r["ag"] else ("X" if r["hg"] == r["ag"] else "2")
            models = {
                "multipliers_1.10_0.90": (max(.05,1.10*bh),max(.05,.90*ba)),
                "neutral_1.00_1.00": (max(.05,bh),max(.05,ba)),
                "venue_adjusted": (max(.05,(uh[0]+ua[1])/2),max(.05,(ua[0]+uh[1])/2)),
            }
            rec = {k:r[k] for k in ["league","season","date","home","away","hg","ag","source_url","source_sha256"]}
            rec["actual_1X2"] = actual
            for name,(lh,la) in models.items():
                p = probs_1x2(lh,la); labels = ["1","X","2"]; idx = labels.index(actual)
                pred = labels[int(np.argmax(p))]
                rec.update({f"{name}_lambda_home":lh,f"{name}_lambda_away":la,
                    f"{name}_p1":p[0],f"{name}_px":p[1],f"{name}_p2":p[2],
                    f"{name}_prediction":pred,f"{name}_correct":int(pred==actual),
                    f"{name}_brier":sum((p[i]-(1 if i==idx else 0))**2 for i in range(3)),
                    f"{name}_logloss":-math.log(max(p[idx],1e-15))})
            detail.append(rec)
            history.extend(events(r))
    return pd.DataFrame(detail),pd.DataFrame(skipped)

def metrics(detail):
    names = ["multipliers_1.10_0.90","neutral_1.00_1.00","venue_adjusted"]
    glob, group_rows = [], []
    for name in names:
        if detail.empty: continue
        glob.append({"Modelo":name,"Partidos":len(detail),"Aciertos":int(detail[f"{name}_correct"].sum()),
            "Exactitud":float(detail[f"{name}_correct"].mean()),"Brier":float(detail[f"{name}_brier"].mean()),
            "LogLoss":float(detail[f"{name}_logloss"].mean()),
            "Nota":"Walk-forward exploratorio; no mide ROI/EV"})
    for (league,season), g in detail.groupby(["league","season"]):
        for name in names:
            group_rows.append({"Liga":league,"Temporada":season,"Modelo":name,"Partidos":len(g),
                "Aciertos":int(g[f"{name}_correct"].sum()),"Exactitud":float(g[f"{name}_correct"].mean()),
                "Brier":float(g[f"{name}_brier"].mean()),"LogLoss":float(g[f"{name}_logloss"].mean())})
    return pd.DataFrame(glob),pd.DataFrame(group_rows)

def main():
    raw_rows, manifest, errors = [], [], []
    for league, code in LEAGUES.items():
        for season in SEASONS:
            url = BASE_URL.format(season=season,code=code)
            try:
                rows, m = fetch(league,code,season)
                raw_rows.extend(rows); manifest.append(m)
            except Exception as e:
                errors.append({"league":league,"season":season,"url":url,"status":"ERROR","error":str(e)})
    if not raw_rows:
        raise RuntimeError("No se descargaron fuentes. Compruebe la conexión o coloque los JSON en data/openfootball_raw/.")
    raw = pd.DataFrame(raw_rows)
    cleaned, invalid, nonint, duplicates, conflicts = clean(raw)
    detail, skipped = backtest(cleaned)
    global_metrics, seasonal_metrics = metrics(detail)
    quality = pd.DataFrame([
        ["Ejecución UTC",RUN_UTC,"Trazabilidad"],
        ["Fuentes descargadas",len(manifest),f"Esperadas {len(LEAGUES)*len(SEASONS)}"],
        ["Fuentes fallidas",len(errors),"Ver manifiesto"],
        ["Filas de fixture",len(raw),"Incluye partidos sin marcador final"],
        ["Marcadores incompletos/filas inválidas",len(invalid),"No se usan para el backtest"],
        ["Marcadores no enteros",len(nonint),"Excluidos"],
        ["Duplicados idénticos",len(duplicates),"Registrados; una copia retenida"],
        ["Filas en conflictos de marcador",len(conflicts),"Excluidas de forma explícita"],
        ["Partidos limpios únicos",len(cleaned),"Después de limpieza"],
        ["Partidos evaluables",len(detail),"Mínimo cinco partidos previos por equipo"],
        ["Partidos omitidos",len(skipped),"Motivo por partido en hoja Omitidos"],
        ["ROI/EV","No calculado","No hay cuotas históricas verificadas"],
        ["Corte temporal","Solo partidos con fecha anterior","Partidos del mismo día excluidos del historial"],
    ],columns=["Control","Resultado","Nota"])
    OUT.parent.mkdir(parents=True,exist_ok=True)
    MANIFEST.parent.mkdir(parents=True,exist_ok=True)
    MANIFEST.write_text(json.dumps({"generated_utc":RUN_UTC,"sources":manifest,"errors":errors},ensure_ascii=False,indent=2),encoding="utf-8")
    with pd.ExcelWriter(OUT,engine="openpyxl") as writer:
        pd.DataFrame(manifest+errors).to_excel(writer,index=False,sheet_name="Fuentes")
        global_metrics.to_excel(writer,index=False,sheet_name="Metricas_globales")
        seasonal_metrics.to_excel(writer,index=False,sheet_name="Metricas_liga_temporada")
        detail.to_excel(writer,index=False,sheet_name="Detalle_backtest")
        skipped.to_excel(writer,index=False,sheet_name="Omitidos")
        duplicates.to_excel(writer,index=False,sheet_name="Duplicados")
        conflicts.to_excel(writer,index=False,sheet_name="Conflictos_marcador")
        invalid.to_excel(writer,index=False,sheet_name="Filas_invalidas")
        nonint.to_excel(writer,index=False,sheet_name="Marcadores_no_enteros")
        quality.to_excel(writer,index=False,sheet_name="Controles_calidad")
        raw.to_excel(writer,index=False,sheet_name="Datos_importados")
        for ws in writer.book.worksheets:
            ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
            for c in ws[1]:
                c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E78")
                c.alignment=Alignment(wrap_text=True,vertical="top")
            for col in ws.columns:
                letter=col[0].column_letter
                maxlen=max((len(str(c.value)) if c.value is not None else 0) for c in col[:100])
                ws.column_dimensions[letter].width=min(max(maxlen+2,12),48)
    print(f"Workbook: {OUT}")
    print(f"Fuentes: {len(manifest)}/15; fixtures: {len(raw)}; válidos: {len(cleaned)}; evaluables: {len(detail)}; omitidos: {len(skipped)}")
    print(f"Duplicados exactos: {len(duplicates)}; conflictos: {len(conflicts)}")
    if errors: print("ADVERTENCIA: ejecución parcial, no es un backtest completo.")

if __name__ == "__main__":
    main()
